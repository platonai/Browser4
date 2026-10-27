# -*- coding: utf-8 -*-
"""Build the final spreadsheet for the Amazon microwave best-seller crawl.

Field order and column names come from
browser4-tests/pulsar-tests-common/src/test/resources/asin_field_description.xlsx (Sheet2 header row).
"""
import json
import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = r"D:\workspace\Browser4\Browser4-4.15-feat"
OUT = os.path.join(ROOT, "output", "amazon-microwaves-2026-10-26")
SRC = os.path.join(ROOT, "browser4-tests", "pulsar-tests-common", "src", "test",
                   "resources", "asin_field_description.xlsx")

BESTSELLER_PAGE = ("https://www.amazon.com/Best-Sellers-Kitchen-Dining-"
                   "Microwave-Ovens/zgbs/kitchen/289935")
LABEL = "microwave-bestsellers-2026-10-26"
KEYWORDS = "microwave ovens"
CATEGORY_ID = "289935"
MAXCELL = 30000

# extra (non-exported) keys collected by the extractor for QA / post-processing
EXTRA_KEYS = ["merchantInfoText", "buyboxScopeText", "couponText", "couponText",
              "brandHref", "sellerHref", "finalUrl", "bestSellerRankPage"]


def read_header():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Sheet2"]
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    names = [c for c in rows[0] if c]
    descs = list(rows[1])[:len(names)] if len(rows) > 1 else []
    legend = [(n, (descs[i] if i < len(descs) else "") or "") for i, n in enumerate(names)]
    return names, legend


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    s = s.replace("\x00", "")
    s = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", " ", s)
    if len(s) > MAXCELL:
        s = s[:MAXCELL] + "...[truncated]"
    return s.strip()


def load_record(asin):
    for sub in ("raw3", "raw2", "raw"):
        p = os.path.join(OUT, sub, asin + ".json")
        if os.path.exists(p):
            try:
                with open(p, encoding="utf-8") as f:
                    return json.load(f), sub
            except Exception:
                continue
    return None, None


def brand_from_overview(rec):
    m = re.search(r"Brand:\s*([^|]+)", rec.get("overviewbullets") or "")
    if m:
        return m.group(1).strip()
    try:
        kv = json.loads(rec.get("detailbullets") or "{}")
        for k in ("Brand Name", "Brand", "Manufacturer"):
            if kv.get(k):
                return kv[k]
    except Exception:
        pass
    return ""


def main():
    header, legend = read_header()
    with open(os.path.join(OUT, "bs_list.json"), encoding="utf-8") as f:
        bs = json.load(f)
    items = bs["items"][:20]

    rows, missing = [], []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for idx, it in enumerate(items, start=1):
        asin = it["asin"]
        d, src = load_record(asin)
        if d is None:
            missing.append(asin)
            continue

        rec = dict(d)
        rec["id"] = idx
        rec["url"] = "https://www.amazon.com/dp/" + asin
        href = it.get("href") or ""
        rec["href"] = ("https://www.amazon.com" + href) if href.startswith("/") else href
        rec["label"] = LABEL
        rec["keywords"] = KEYWORDS
        rec["category"] = CATEGORY_ID
        rec["task_time"] = now
        rec["referer"] = BESTSELLER_PAGE

        # rank: page rank block when present, otherwise the position on today's best-seller list
        if not (rec.get("rank2") or "").strip():
            fallback = "%s in Microwave Ovens (Best Sellers list)" % (it.get("rank") or "#%d" % idx)
            rec["rank"] = fallback
            rec["rank2"] = fallback
        if not (rec.get("rank") or "").strip():
            rec["rank"] = rec.get("rank2", "")

        if not (rec.get("brand") or "").strip():
            rec["brand"] = brand_from_overview(rec)
        if not (rec.get("title") or "").strip():
            rec["title"] = it.get("title", "")
        if not (rec.get("imgsrc") or "").strip():
            rec["imgsrc"] = it.get("listImg", "")
        if not (rec.get("price") or "").strip():
            rec["price"] = it.get("listPrice", "")
        if not (rec.get("buyboxprice") or "").strip():
            rec["buyboxprice"] = rec.get("price") or it.get("listPrice", "")
        if not (rec.get("soldby") or "").strip():
            mi = rec.get("merchantInfoText") or ""
            m = re.search(r"Shipper\s*/\s*Seller\s*[:\s]\s*(.+)", mi)
            if m:
                rec["soldby"] = m.group(1).strip()
        if not (rec.get("shipsfrom") or "").strip():
            rec["shipsfrom"] = rec.get("soldby", "")
        if not (rec.get("sellerID") or "").strip() and "amazon" in (rec.get("soldby") or "").lower():
            rec["sellerID"] = "ATVPDKIKX0DER"
        if not (rec.get("feedbackurl") or "").strip() and rec.get("sellerID"):
            rec["feedbackurl"] = "https://www.amazon.com/sp?seller=" + rec["sellerID"]

        rows.append(rec)

    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = "microwaves_top20"

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="305496")
    for c, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, rec in enumerate(rows, start=2):
        for c, name in enumerate(header, start=1):
            ws.cell(row=r, column=c, value=clean(rec.get(name, "")))

    widths = {"id": 5, "title": 55, "url": 40, "href": 30, "asin": 12, "pasin": 12,
              "price": 10, "buyboxprice": 11, "listprice": 10, "yousave": 14,
              "score": 7, "reviews": 9, "rank2": 38, "brand": 22, "soldby": 16,
              "shipsfrom": 14, "weight": 14, "volume": 26, "categoryname": 24,
              "instock": 12, "stock_status": 11, "label": 22, "keywords": 16,
              "task_time": 20, "deliverto": 28, "withdeal": 16, "othersellernum": 10,
              "score5percent": 9, "score4percent": 9, "score3percent": 9,
              "score2percent": 9, "score1percent": 9, "gallery": 30, "pay": 10}
    for c, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), len(rows) + 1)

    # legend sheet: field name + its description from the source workbook
    lg = out_wb.create_sheet("fields")
    lg.cell(row=1, column=1, value="field").font = hdr_font
    lg.cell(row=1, column=1).fill = hdr_fill
    lg.cell(row=1, column=2, value="description").font = hdr_font
    lg.cell(row=1, column=2).fill = hdr_fill
    for r, (nm, ds) in enumerate(legend, start=2):
        lg.cell(row=r, column=1, value=str(nm))
        lg.cell(row=r, column=2, value=str(ds))
    lg.column_dimensions["A"].width = 24
    lg.column_dimensions["B"].width = 42
    lg.freeze_panes = "A2"

    dest = os.path.join(OUT, "amazon_microwaves_top20.xlsx")
    out_wb.save(dest)

    # QA summary
    print("saved:", dest)
    print("rows:", len(rows), "cols:", len(header), "missing:", missing)
    for name in header:
        filled = sum(1 for r in rows if str(r.get(name, "")).strip() not in ("", "0", "None")
                     or str(r.get(name, "")).strip() == "0" and name.startswith("is"))
        print("  %-22s %2d/%d" % (name, filled, len(rows)))


if __name__ == "__main__":
    main()


